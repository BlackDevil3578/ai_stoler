import streamlit as st
import pandas as pd
from PyPDF2 import PdfReader
import xlsxwriter
from tempfile import NamedTemporaryFile
import numpy as np
import os
import matplotlib.pyplot as plt
from collections import Counter
import io
import nltk
import pdfplumber
import re
import string
from nltk.tokenize import word_tokenize
from nltk.probability import FreqDist
from docx import Document

# Download NLTK resources only once
@st.cache_resource
def download_nltk_resources():
    nltk.download('punkt')
    nltk.download('stopwords')

# Functions for DataFrame manipulation

def save_text_to_word(text_content, output_path):
    doc = Document()
    doc.add_paragraph(text_content)
    doc.save(output_path)



def convert_excel_to_sql_file(df, table_name, output_path):
    sql_content = f"CREATE TABLE {table_name} (\n\tPage INT PRIMARY KEY,\n\tContent TEXT\n);\n\n"
    for _, row in df.iterrows():
        sql_content += f"INSERT INTO {table_name} (Page, Content) VALUES ({row['Page']}, '{row['Content']}');\n"
    with open(output_path, 'w', encoding='utf-8') as file:
        file.write(sql_content)

def extract_text_from_pdf(pdf_reader, numerate_pages):
    all_page_content = ""
    for page_number in range(len(pdf_reader.pages)):
        page = pdf_reader.pages[page_number]
        page_content = page.extract_text()
        if not page_content:
            page_content = ""

        if numerate_pages:
            all_page_content += f"Page {page_number + 1}:\n{page_content if page_content else ''}\n\n"
        else:
            all_page_content += f"{page_content if page_content else ''}\n\n"
    return all_page_content

def split_text_by_words(text_content, start_word, end_word):
    # Split text content by non-word boundaries and word boundaries
    words = re.split(r'(\W+)', text_content)
    start_word = start_word.lower()
    end_word = end_word.lower()

    segment = []
    capturing = False
    content_segments = []
    current_segment = ""

    for word in words:
        # Normalize the word to lower case to match start and end words case-insensitively
        word_normalized = word.lower()

        if capturing:
            current_segment += word
            if word_normalized == end_word:
                capturing = False
                content_segments.append(current_segment.strip())
                current_segment = ""

        elif word_normalized == start_word:
            capturing = True
            current_segment = word

    if current_segment:
        content_segments.append(current_segment.strip())

    return content_segments


# Streamlit app configuration
def configure_streamlit():
    st.set_page_config(
        page_title="PDF-stoler",
        page_icon="C:\\Users\\Aymen\\PycharmProjects\\pdf to excel and text\\icon2.png",
        layout="wide",
        initial_sidebar_state="expanded"
    )

    with open("C:\\Users\\Aymen\\PycharmProjects\\pdf to excel and text\\style.css") as f:
        st.markdown(f"<style>{f.read()}</style>", unsafe_allow_html=True)

    st.markdown("<h1 class='title'>Upload your pdf file here :</h1>", unsafe_allow_html=True)

def handle_file_upload():

    uploaded_file = st.file_uploader("Choose a PDF file", type="pdf")
    return uploaded_file

def process_uploaded_file(uploaded_file):
    table_name_raw = os.path.splitext(os.path.basename(uploaded_file.name))[0]
    table_name = re.sub(r'\W+', '_', table_name_raw)

    with st.expander("Would you like to include page numbers in the text content? If yes, please check the box", expanded=True):
        if 'numerate_pages' not in st.session_state:
            st.session_state.numerate_pages = False
        numerate_pages = st.checkbox("Include page numbers in the text content", value=st.session_state.numerate_pages)
        st.session_state.numerate_pages = numerate_pages

    with st.expander("Would you like to split the PDF by specific words? If yes, please check the box", expanded=True):
        if 'custom_split' not in st.session_state:
            st.session_state.custom_split = False
        custom_split = st.checkbox("Split PDF by specific words", value=st.session_state.custom_split)
        st.session_state.custom_split = custom_split

        if custom_split:
            if 'start_word' not in st.session_state:
                st.session_state.start_word = ''
            if 'end_word' not in st.session_state:
                st.session_state.end_word = ''

            start_word = st.text_input("Start word", key='start_word')
            end_word = st.text_input("End word", key='end_word')
            split_button_clicked = st.button("Split PDF by Words")

            if start_word and end_word:
                # Do not reassign session state variables here
                if not split_button_clicked:
                    split_button_clicked = True
        else:
            start_word = ""
            end_word = ""
            split_button_clicked = False

    pdf_reader = PdfReader(uploaded_file)
    return pdf_reader, numerate_pages, custom_split, start_word, end_word, table_name, split_button_clicked

def extract_content_and_save_to_excel(pdf_reader, numerate_pages, custom_split, start_word, end_word, split_button_clicked):
    with NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp:
        temp_excel_path = tmp.name

    workbook = xlsxwriter.Workbook(temp_excel_path)
    worksheet = workbook.add_worksheet("First worksheet")
    worksheet.write(0, 0, "Page")
    worksheet.write(0, 1, "Content")

    line = 1
    all_page_content = ""
    content_segments = []

    for page_number in range(len(pdf_reader.pages)):
        page = pdf_reader.pages[page_number]
        page_content = page.extract_text()
        if not page_content:
            page_content = ""

        worksheet.write(line, 0, page_number + 1)
        worksheet.write(line, 1, page_content)
        line += 1

        if numerate_pages:
            all_page_content += f"Page {page_number + 1}:\n{page_content if page_content else ''}\n\n"
        else:
            all_page_content += f"{page_content if page_content else ''}\n\n"

    if custom_split and start_word and end_word and split_button_clicked:
        all_page_content_without_pagenumbers = re.sub(r"Page \d+:\n", "", all_page_content)
        content_segments = split_text_by_words(all_page_content_without_pagenumbers, start_word, end_word)

        workbook = xlsxwriter.Workbook(temp_excel_path)
        worksheet = workbook.add_worksheet("First worksheet")
        worksheet.write(0, 0, "Page")
        worksheet.write(0, 1, "Content")

        for i, content in enumerate(content_segments):
            worksheet.write(i + 1, 0, i + 1)
            worksheet.write(i + 1, 1, content)

        workbook.close()
    else:
        workbook.close()

    return temp_excel_path, all_page_content


def load_dataframe(temp_excel_path):
    if 'df' not in st.session_state:
        st.session_state.df = pd.read_excel(temp_excel_path, index_col=None)
    else:
        st.session_state.df = pd.read_excel(temp_excel_path, index_col=None)

def display_dataframes():
    st.markdown("<h2 class='title'>Data from Excel file:</h2>", unsafe_allow_html=True)
    # Display the dataframe without providing options for adding/modifying/deleting rows
    df_display = st.empty()
    df_display.dataframe(st.session_state.df, height=500)
    return df_display


def save_text_content(all_page_content):
    with open('TextFile.txt', 'w', encoding='utf-8') as f:
        f.write(all_page_content)
    st.session_state.text_content = all_page_content

def toggle_columns():
    if 'show_left_col' not in st.session_state:
        st.session_state.show_left_col = True
    if 'show_right_col' not in st.session_state:
        st.session_state.show_right_col = True

    toggle_left_button_text = "Toggle DataFrame"
    toggle_right_button_text = "Toggle Text Content"

    col1, col2 = st.columns(2)

    with col1:
        if st.button(toggle_left_button_text):
            st.session_state.show_left_col = not st.session_state.show_left_col

    with col2:
        if st.button(toggle_right_button_text):
            st.session_state.show_right_col = not st.session_state.show_right_col

    return st.session_state.show_left_col, st.session_state.show_right_col


def display_dataframes():
    st.markdown("<h2 class='title'>Data from Excel file:</h2>", unsafe_allow_html=True)
    df_display = st.empty()
    df_display.dataframe(st.session_state.df, height=500)
    return df_display


def save_df_to_excel(df, temp_excel_path):
    with pd.ExcelWriter(temp_excel_path, engine='xlsxwriter') as writer:
        df.to_excel(writer, index=False, sheet_name='First worksheet')




def extract_text_from_pdf_plumber(file_path):
    with pdfplumber.open(file_path) as pdf:
        pdf_text = ""
        for page in pdf.pages:
            pdf_text += page.extract_text()
    return pdf_text

# Preprocess the text
def preprocess_text(text):
    text = text.lower()  # Convert to lowercase
    text = re.sub(r'\d+', '', text)  # Remove digits
    text = text.translate(str.maketrans('', '', string.punctuation))  # Remove punctuation
    text = text.strip()  # Remove leading/trailing whitespace
    return text

# Function to count total words
def count_total_words(text):
    words = word_tokenize(preprocess_text(text))
    words = [word for word in words if word.isalnum()]  # Filter out non-alphanumeric tokens
    return len(words)

# Function to count specific word occurrences using NLTK
def count_specific_word(text, word):
    words = word_tokenize(preprocess_text(text))
    words = [w for w in words if w.isalnum()]  # Filter out non-alphanumeric tokens
    word_freq = FreqDist(words)
    return word_freq[word]

# Update the Streamlit part for counting words
def display_text_content():
    download_nltk_resources()
    st.markdown("<h2 class='title'>Text Content from PDF:</h2>", unsafe_allow_html=True)
    st.text_area("", st.session_state.text_content, height=500)

    with st.expander("Total words count", expanded=True):
        # Word Count
        if st.button("Count Total Words"):
            total_words = count_total_words(st.session_state.text_content)
            st.write(f"Total words: {total_words}")

    # Count Specific Word
    with st.expander("Count Specific Word occurrence", expanded=True):
        word_to_count = st.text_input("Enter word to count")
        if st.button("Count Specific Word"):
            word_count = count_specific_word(st.session_state.text_content, word_to_count)
            st.write(f"The word '{word_to_count}' occurred {word_count} times.")

    # Word Frequency Analysis
    with st.expander("Show Top 10 Words Frequency", expanded=True):
        if st.button("Show Top 10 Words Frequency"):
            try:
                words = word_tokenize(preprocess_text(st.session_state.text_content))
                words = [word for word in words if word.isalnum()]  # Filter out non-alphanumeric tokens
                word_freq = Counter(words)
                most_common_words = word_freq.most_common(10)
                words_list, frequencies = zip(*most_common_words)

                colors = plt.cm.Blues(np.linspace(0.7, 0.3, len(words_list)))

                fig, ax = plt.subplots(figsize=(24, 13.5), dpi=300)
                bars = ax.bar(words_list, frequencies, color=colors)
                ax.set_xlabel('Words', fontsize=20, color='navy')
                ax.set_ylabel('Frequency', fontsize=20, color='navy')
                ax.set_title('Top 10 Word Frequency Distribution', fontsize=24, color='navy')
                plt.xticks(rotation=45, ha='right', fontsize=15, color='navy')
                plt.yticks(fontsize=15, color='navy')

                for bar in bars:
                    yval = bar.get_height()
                    plt.text(bar.get_x() + bar.get_width() / 2, yval + 0.5, int(yval), ha='center', fontsize=15, color='navy')

                buf = io.BytesIO()
                plt.savefig(buf, format="png", bbox_inches='tight', dpi=300)
                buf.seek(0)

                st.image(buf)
                st.download_button(
                    label="Download Plot as PNG",
                    data=buf,
                    file_name="top_10_word_frequency.png",
                    mime="image/png",
                    help="Click to download the plot as a high-quality PNG file"
                )
                buf.close()

            except Exception as e:
                st.error(f"An error occurred: {e}")

def add_download_buttons(temp_excel_path, pdf_filename):
    st.sidebar.title("Download your files:")

    with open(temp_excel_path, "rb") as file:
        st.sidebar.download_button(
            label="Download Excel File",
            data=file,
            file_name=f"{pdf_filename}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            key='xlsx-download',
            help="Click to download the Excel file"
        )

    with open("TextFile.txt", "rb") as file:
        st.sidebar.download_button(
            label="Download Text File",
            data=file,
            file_name=f"{pdf_filename}.txt",
            mime="text/plain",
            key='txt-download',
            help="Click to download the text file"
        )

    with NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
        temp_word_path = tmp.name

    save_text_to_word(st.session_state.text_content, temp_word_path)

    with open(temp_word_path, "rb") as file:
        st.sidebar.download_button(
            label="Download Word File",
            data=file,
            file_name=f"{pdf_filename}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            help="Click to download the Word file"
        )

def convert_and_download_sql(table_name, pdf_filename):
    st.sidebar.title("Convert and Download SQL Database:")
    if st.sidebar.button("Convert to SQL Database"):
        with NamedTemporaryFile(delete=False, suffix=".sql") as tmp:
            temp_sql_path = tmp.name

        convert_excel_to_sql_file(st.session_state.df, table_name, temp_sql_path)

        with open(temp_sql_path, "rb") as file:
            st.sidebar.download_button(
                label="Download SQL File",
                data=file,
                file_name=f"{pdf_filename}.sql",
                mime="application/sql",
                help="Click to download the SQL file"
            )

# Main function
def main():
    configure_streamlit()
    uploaded_file = handle_file_upload()

    if uploaded_file is not None:
        pdf_filename = os.path.splitext(uploaded_file.name)[0]
        pdf_reader, numerate_pages, custom_split, start_word, end_word, table_name, split_button_clicked = process_uploaded_file(
            uploaded_file)
        temp_excel_path, all_page_content = extract_content_and_save_to_excel(pdf_reader, numerate_pages, custom_split,
                                                                              start_word, end_word,
                                                                              split_button_clicked)
        load_dataframe(temp_excel_path)
        save_text_content(all_page_content)

        show_left_col, show_right_col = toggle_columns()

        # Show the DataFrame in the left column if toggled
        if show_left_col or show_right_col:
            col1, col2 = st.columns(2)
            if show_left_col:
                with col1:
                    display_dataframes()

            if show_right_col:
                with col2:
                    display_text_content()

        add_download_buttons(temp_excel_path, pdf_filename)
        convert_and_download_sql(table_name, pdf_filename)

if __name__ == "__main__":
    main()
